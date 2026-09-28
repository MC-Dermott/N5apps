"""Completed-spreadsheet solutions for download/fill/upload spreadsheet questions.

A spreadsheet question's generator builds its completed workbook (live formulas in the
cells pupils fill, same layout as the question workbook) and attaches it with
`solution_metadata()`:

    metadata={
        "spreadsheet_bytes": ...,                       # the blank question workbook
        "spreadsheet_filename": "savings_schedule.xlsx",
        "spreadsheet_answer_cell": ("Savings", "B26"),
        **solution_metadata(wb, ws, values, filled, min_row=13, max_row=26),
    }

`core/ui/solution_ui.py`'s `render_solution()` then shows it in the worked solution as a
download plus a grid of the completed values (and the formulas, where there are any).
openpyxl can't evaluate formulas, so the generator passes the values its own Python
computed for every formula cell in `values`.
"""
import io
from datetime import date, datetime

from openpyxl.utils import get_column_letter


def _format(value, number_format):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = number_format or "General"
        decimals = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
        if fmt.endswith("%"):
            return f"{value * 100:.{decimals}f}%"
        if "£" in fmt:
            return f"£{value:,.{decimals}f}"
        if fmt != "General":
            return f"{value:,.{decimals}f}"
        return f"{value:g}" if isinstance(value, float) else str(value)
    return str(value)


def solution_view(ws, values, filled, min_row, max_row, min_col=1, max_col=None):
    """A render-ready snapshot of rows min_row..max_row of the completed sheet `ws`.

    `values` maps a cell ref to the value its formula evaluates to; `filled` is every cell
    ref the pupil had to complete (highlighted in the rendered grid)."""
    max_col = max_col or ws.max_column
    columns = [get_column_letter(c) for c in range(min_col, max_col + 1)]
    rows, formulas = [], {}
    for r in range(min_row, max_row + 1):
        shown = []
        for letter in columns:
            ref = f"{letter}{r}"
            cell = ws[ref]
            raw = cell.value
            if isinstance(raw, str) and raw.startswith("="):
                formulas[ref] = raw
                raw = values[ref]
            shown.append(_format(raw, cell.number_format))
        rows.append((r, shown))
    return {"sheet": ws.title, "columns": columns, "rows": rows,
            "filled": sorted(filled), "formulas": formulas}


def workbook_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def solution_metadata(wb, ws, values, filled, min_row, max_row, min_col=1, max_col=None,
                      filename=None):
    meta = {
        "spreadsheet_solution_bytes": workbook_bytes(wb),
        "spreadsheet_solution_view": solution_view(ws, values, filled, min_row, max_row,
                                                   min_col, max_col),
    }
    if filename:
        meta["spreadsheet_solution_filename"] = filename
    return meta
