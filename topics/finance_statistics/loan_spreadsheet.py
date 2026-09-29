"""Downloadable loan-schedule spreadsheets (download, fill in, upload) — shared by the
Higher Finance topics whose worksheets build a loan schedule in Excel (Loan Schedules,
Mortgages …).

The layout matches the sheets in Higher Apps/Worksheets/Finance/Questions/
Loan_Schedules_Worksheet.xlsx cell for cell, so the worksheet's cell references (C6, D13 …)
work here too:

    B3 Price   B4 Loan amount   B5 Annual effective rate   B6 Monthly effective rate
    B7 Level monthly repayment  B8 Final repayment         B9 Term (months)
    row 11 headers; row 12 = month 0; rows 13… = months 1…n

    D13 = ROUND($C$6*F12,2)   E13 = C13-D13   F13 = F12-E13   (the last month repays $C$8)

Two modes, matching the worksheet's two Goal Seek skills:
- "repayment": rate given; pupils Goal Seek the level repayment (graded cell C7).
- "rate": repayments given; pupils Goal Seek the monthly rate, then convert it to the
  annual effective rate (graded cell C5, as a percentage).
"""
import io
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.engine.spreadsheet_solution import excel_round, solution_metadata

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
ANSWER_FILL = PatternFill("solid", fgColor="FFFF00")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MONEY = "#,##0.00"
RATE = "0.000%"
FIRST_ROW = 13


def monthly_rate(annual):
    return (1 + annual) ** (1 / 12) - 1


def level_repayment(P, i, n):
    return P * i / (1 - (1 + i) ** -n)


def schedule(P, i, R, n, final=None):
    """Rows (month, repayment, interest, capital, outstanding) with the SQA rounding: interest
    to the penny on the previous outstanding. If `final` is None the last repayment is
    whatever clears the loan (outstanding + that month's interest).

    Mirrors the sheet's arithmetic exactly — capital (C−D) and outstanding (F−E) are never
    rounded, so the float drift a pupil's spreadsheet accumulates is reproduced and ROUND()
    breaks the same way. Values are formatted to the penny only when displayed."""
    bal, rows = P, []
    for t in range(1, n + 1):
        interest = excel_round(i * bal)
        if t == n:
            pay = (bal + interest) if final is None else final
        else:
            pay = R
        capital = pay - interest
        bal = bal - capital
        rows.append((t, pay, interest, capital, bal))
    return rows


def clear_schedule(P, i, R, max_months=600):
    """Pay R a month until the loan is cleared; the last repayment is only what is left
    (outstanding + that month's interest). Same float-faithful arithmetic as schedule()."""
    bal, rows = P, []
    while bal > 1e-9 and len(rows) < max_months:
        interest = excel_round(i * bal)
        pay = min(R, bal + interest)
        capital = pay - interest
        bal = bal - capital
        rows.append((len(rows) + 1, pay, interest, capital, bal))
    return rows


def target_repayment(P, i, n, target):
    """The smallest repayment, to the penny, that leaves no more than `target` outstanding
    after n months (Goal Seek to the target, then round UP)."""
    g = (1 + i) ** n
    R = round(-(-(P * g - target) * i / (g - 1) * 100 // 1)) / 100      # continuous answer, rounded up
    while schedule(P, i, round(R - 0.01, 2), n, round(R - 0.01, 2))[-1][4] <= target:
        R = round(R - 0.01, 2)
    while schedule(P, i, R, n, R)[-1][4] > target:
        R = round(R + 0.01, 2)
    return R


LOAN_LABELS = {
    "rows": [(3, "Price (£)"), (4, "Loan amount (£)"), (5, "Annual effective rate"), (6, "Monthly effective rate"),
             (7, "Level monthly repayment (£)"), (8, "Final repayment (£)"), (9, "Term (months)")],
    "outstanding": "Loan outstanding (£)",
}
MORTGAGE_LABELS = {
    "rows": [(4, "Mortgage amount (£)"), (5, "Annual effective rate"), (6, "Monthly effective rate"),
             (7, "Level monthly repayment (£)"), (8, "Final repayment (£)"), (9, "Term (months)"),
             (10, "Target mortgage outstanding (£)")],
    "outstanding": "Mortgage outstanding (£)",
}


def schedule_sheet(title, sheet_name, n, labels=LOAN_LABELS):
    """A blank schedule sheet in the worksheet's layout (see the module docstring), with month
    numbers 0…n in column B. Returns (wb, ws)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    for col in "CDEF":
        ws.column_dimensions[col].width = 18
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    for row, label in labels["rows"]:
        ws[f"B{row}"] = label
    for ref, fmt in [("C3", MONEY), ("C4", MONEY), ("C5", RATE), ("C6", RATE), ("C7", MONEY), ("C8", MONEY),
                     ("C10", MONEY)]:
        ws[ref].number_format = fmt
    ws["C9"] = n
    for col, text in zip("BCDEF", ["Time (months)", "Repayment (£)", "Interest content of repayment (£)",
                                   "Capital content of repayment (£)", labels["outstanding"]]):
        cell = ws[f"{col}11"]
        cell.value = text
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[11].height = 45
    ws["B12"] = 0
    ws["F12"].number_format = MONEY
    for t in range(1, n + 1):
        r = FIRST_ROW + t - 1
        ws[f"B{r}"] = t
        for col in "CDEF":
            ws[f"{col}{r}"].number_format = MONEY
    return wb, ws


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_loan_spreadsheet(*, title, sheet_name, filename, mode, P, i, n, rows, R, final,
                           annual=None, price=None, deposit=None):
    """Question workbook bytes, graded cell, and the completed-solution metadata.

    `rows` is `schedule(P, i, R, n, final)`; `annual` is the annual effective rate (a
    fraction); `price`/`deposit` (a fraction) for a finance deal, where the loan is the
    price minus the deposit."""
    last = FIRST_ROW + n - 1
    deal = price is not None

    # ---------------- question ----------------
    wb, ws = schedule_sheet(title, sheet_name, n)
    to_fill = []
    if deal:
        ws["C3"] = price
        ws["D3"] = f"Deposit: {deposit:.0%}"
        to_fill.append("C4")
    else:
        ws["C3"] = "—"
        ws["C4"] = P
        ws["F12"] = P
    if mode == "repayment":
        ws["C5"] = annual
        to_fill += ["C6", "C7", "C8"]
    else:
        ws["C7"], ws["C8"] = R, final
        to_fill += ["C5", "C6"]
    if deal:
        to_fill.append("F12")
    to_fill += [f"{col}{r}" for r in range(FIRST_ROW, last + 1) for col in "CDEF"]
    for ref in to_fill:
        ws[ref].fill = ANSWER_FILL
    ws["A2"] = "Fill in the yellow cells, save the file, then upload it."
    ws["A2"].font = Font(italic=True)
    answer_cell = "C7" if mode == "repayment" else "C5"
    question_bytes = _bytes(wb)

    # ---------------- completed solution ----------------
    wb, ws = schedule_sheet(title, sheet_name, n)
    ws["A2"] = "Completed solution: the yellow cells show the formulas used."
    ws["A2"].font = Font(italic=True)
    values = {}
    if deal:
        ws["C3"] = price
        ws["D3"] = f"Deposit: {deposit:.0%}"
        ws["C4"] = f"=C3*(1-{deposit:g})"
        values["C4"] = P
    else:
        ws["C3"] = "—"
        ws["C4"] = P
    if mode == "repayment":
        ws["C5"] = annual
        ws["C6"] = "=(1+C5)^(1/12)-1"
        ws["C7"] = R
        ws["D7"] = "← Goal Seek: last 'Loan outstanding' = 0 by changing C7, then round to 2 d.p."
        ws["C8"] = f"=F{last - 1}+ROUND($C$6*F{last - 1},2)"
        values.update({"C6": i, "C8": final})
    else:
        ws["C5"] = "=(1+C6)^12-1"
        ws["C6"] = i
        ws["D6"] = "← Goal Seek: set the last 'Loan outstanding' to 0 by changing C6"
        ws["C7"], ws["C8"] = R, final
        values["C5"] = annual
    ws["F12"] = "=C4"
    values["F12"] = P
    for t, pay, interest, capital, bal in rows:
        r = FIRST_ROW + t - 1
        ws[f"C{r}"] = "=$C$8" if t == n else "=$C$7"
        ws[f"D{r}"] = f"=ROUND($C$6*F{r - 1},2)"
        ws[f"E{r}"] = f"=C{r}-D{r}"
        ws[f"F{r}"] = f"=F{r - 1}-E{r}"
        values.update({f"C{r}": pay, f"D{r}": interest, f"E{r}": capital, f"F{r}": bal})
    for ref in to_fill:
        ws[ref].fill = ANSWER_FILL
    solution = solution_metadata(wb, ws, values, to_fill, min_row=3, max_row=last, min_col=2, max_col=6,
                                 filename=filename.replace(".xlsx", "_solution.xlsx"))

    metadata = {
        "spreadsheet_bytes": question_bytes,
        "spreadsheet_filename": filename,
        "spreadsheet_answer_cell": (sheet_name, answer_cell),
        **solution,
    }
    if mode == "rate":
        metadata["spreadsheet_answer_percent"] = True
    return metadata


# ---------------------------------------------------------------------------
# A change to the plan part-way through — an increased repayment, a payment holiday, a change of
# interest rate — shared by Loan Schedules and Mortgages. Same layout, with the change's details in
# H4:I8. Each *_case(p) takes a plan from level_plan() and returns the part-question's text, answer,
# scaffold, worked solution and spreadsheet metadata, so a topic can ask it on its own or as one part
# of an exam-style question.
# ---------------------------------------------------------------------------

MORTGAGE_SHEET = {"word": "mortgage", "sheet": "Mortgage", "labels": MORTGAGE_LABELS, "file": "mortgage_schedule"}
LOAN_SHEET = {"word": "loan", "sheet": "Loan Schedule", "labels": LOAN_LABELS, "file": "loan_schedule"}


def run_schedule(P, terms, n=None, max_months=600):
    """Schedule rows (month, repayment, interest, capital, outstanding) where terms(t) gives month t's
    (monthly rate, repayment). Without n, runs until cleared and the last repayment is only what is
    left; with n, month n repays whatever is left (the level-repayment final row). Same float-faithful
    arithmetic as schedule()."""
    bal, rows = P, []
    for t in range(1, max_months + 1):
        i, R = terms(t)
        interest = excel_round(i * bal)
        pay = bal + interest if (t == n or (n is None and R >= bal + interest)) else R
        bal = bal - (pay - interest)
        rows.append((t, pay, interest, pay - interest, bal))
        if t == n or (n is None and bal <= 1e-9):
            break
    return rows


def goal_seek_repayment(B, i, m):
    """The repayment (to 2 d.p.) that Goal Seek finds to clear balance B in m months at monthly rate i."""
    def end(R):
        bal = B
        for _ in range(m):
            bal -= R - excel_round(i * bal)
        return bal
    lo, hi = 0.0, B * 2
    for _ in range(80):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if end(mid) > 0 else (lo, mid)
    return excel_round((lo + hi) / 2)


def level_plan(name, P, a, n, kind, title, **extra):
    """A level-repayment plan: R to the penny for months 1…n−1 and the final repayment F."""
    i = monthly_rate(a / 100)
    R = round(level_repayment(P, i, n), 2)
    F = round(schedule(P, i, R, n)[-1][1], 2)
    return dict(name=name, P=P, a=a, i=i, n=n, R=R, F=F, orig=round((n - 1) * R + F, 2), kind=kind,
                title=title, **extra)


def change_sheet(p, subtitle, rows, *, given, extras, fills, pay_ref, rate_ref, answer, grid_rows,
                 c10=None, blank=()):
    """Question workbook bytes, graded cell and completed-solution metadata for a plan that changes
    part-way. `given` — C-column values shown to pupils; `extras` — (row, label, value, format) given in
    H/I; `fills` — (row, label, formula_or_value, value, format) pupils complete in H/I; pay_ref(t) /
    rate_ref(t) — month t's repayment formula and rate cell; `c10` — (label, formula, value)."""
    kind, P, a = p["kind"], p["P"], p["a"]
    n = len(rows)
    last = FIRST_ROW + n - 1

    def base(months):
        wb, ws = schedule_sheet(f"{p['title']} — {subtitle}", kind["sheet"], months, kind["labels"])
        ws.column_dimensions["H"].width = 36
        ws.column_dimensions["I"].width = 16
        if kind["labels"] is LOAN_LABELS:
            ws["C3"] = "—"
        ws["C4"], ws["C5"] = P, a / 100
        for ref, value in given.items():
            ws[ref] = value
        for ref in blank:
            ws[ref] = None
        for row, label, value, fmt in extras:
            ws[f"H{row}"], ws[f"I{row}"] = label, value
            ws[f"I{row}"].number_format = fmt
        for row, label, _, _, fmt in fills:
            ws[f"H{row}"] = label
            ws[f"I{row}"].number_format = fmt
        if c10:
            ws["B10"] = c10[0]
            ws["C10"].number_format = MONEY
        return wb, ws

    to_fill = ["C6", "F12"] + [f"I{row}" for row, *_ in fills] + (["C10"] if c10 else [])
    wb, ws = base(grid_rows)
    for ref in to_fill + [f"{c}{r}" for r in range(FIRST_ROW, FIRST_ROW + grid_rows) for c in "CDEF"]:
        ws[ref].fill = ANSWER_FILL
    ws["A2"] = "Fill in the yellow cells, save the file, then upload it."
    ws["A2"].font = Font(italic=True)
    question_bytes = _bytes(wb)

    wb, ws = base(n)
    ws["A2"] = "Completed solution: the yellow cells show the formulas used."
    ws["A2"].font = Font(italic=True)
    ws["C6"], ws["F12"] = "=(1+C5)^(1/12)-1", "=C4"
    values = {"C6": monthly_rate(a / 100), "F12": P}
    for row, _, content, value, _ in fills:
        ws[f"I{row}"] = content
        values[f"I{row}"] = value
    for t, pay, interest, capital, bal in rows:
        r = FIRST_ROW + t - 1
        rate = rate_ref(t)
        ws[f"C{r}"] = f"=F{r - 1}+ROUND({rate}*F{r - 1},2)" if t == n else pay_ref(t)
        ws[f"D{r}"] = f"=ROUND({rate}*F{r - 1},2)"
        ws[f"E{r}"] = f"=C{r}-D{r}"
        ws[f"F{r}"] = f"=F{r - 1}-E{r}"
        values.update({f"C{r}": pay, f"D{r}": interest, f"E{r}": capital, f"F{r}": bal})
    if c10:
        ws["C10"] = c10[1]
        values["C10"] = c10[2]
    filled = to_fill + [f"{c}{r}" for r in range(FIRST_ROW, last + 1) for c in "CDEF"]
    for ref in filled:
        ws[ref].fill = ANSWER_FILL
    meta = solution_metadata(wb, ws, values, filled, min_row=3, max_row=last, min_col=2, max_col=9,
                             filename=f"{kind['file']}_solution.xlsx")
    return {"spreadsheet_bytes": question_bytes, "spreadsheet_filename": f"{kind['file']}.xlsx",
            "spreadsheet_answer_cell": (kind["sheet"], answer), **meta}


def _change_month(n):
    return random.choice([m for m in (6, 12, 18, 24, 30, 36) if n // 4 <= m <= n - 12])


def increase_case(p):
    P, i, n, R, F, orig, name, w = p["P"], p["i"], p["n"], p["R"], p["F"], p["orig"], p["name"], p["kind"]["word"]
    k = _change_month(n)
    inc = max(10, round(R * random.uniform(0.1, 0.35) / 25) * 25)
    R2 = round(R + inc, 2)
    rows = run_schedule(P, lambda t: (i, R if t <= k else R2))
    last_pay, months = round(rows[-1][1], 2), len(rows)
    total = sum(r[1] for r in rows)
    saving = round(orig - total, 2)
    text = (f"After {k} repayments, {name} increases the monthly repayment by £{inc} to £{R2:,.2f}, starting with "
            f"repayment {k + 1}, and keeps paying this until the {w} is cleared.\n\nDownload the spreadsheet "
            f"below. Complete the schedule to find the final repayment (I7) and the number of months (I8), then "
            f"determine how much money this saves over the term of the {w} (C10). Save the file and upload it.")
    steps = [{"prompt": "Total repaid on the original plan (£)", "answer": orig},
             {"prompt": f"{w.capitalize()} outstanding at the end of month {k} (£)", "answer": round(rows[k - 1][4], 2)},
             {"prompt": f"In which month does the {w} reach £0 paying £{R2:,.2f}?", "answer": months},
             {"prompt": "Final repayment (£)", "answer": last_pay},
             {"prompt": "Saving over the term (£)", "answer": saving}]
    worked = [f"Months 1–{k}: repayment =$C$7 (£{R:,.2f}); from month {k + 1}: =$I$5 (£{R2:,.2f})",
              f"Outstanding at the end of month {k}: £{rows[k - 1][4]:,.2f}",
              f"Fill down: the {w} is cleared in month {months}; the last repayment is only what is left: "
              f"£{last_pay:,.2f}",
              f"Total repaid originally = {n - 1} × {R:,.2f} + {F:,.2f} = £{orig:,.2f}",
              f"Total repaid now = {k} × {R:,.2f} + {months - k - 1} × {R2:,.2f} + {last_pay:,.2f} = £{total:,.2f}",
              f"Saving = £{saving:,.2f}"]
    last = FIRST_ROW + months - 1
    meta = change_sheet(
        p, "increased repayment", rows,
        given={"C7": R, "C8": F, "C9": n},
        extras=[(4, "Repayment increases from month", k + 1, "0"), (5, "New monthly repayment (£)", R2, MONEY)],
        fills=[(7, "Final repayment on the new plan (£)", f"=C{last}", last_pay, MONEY),
               (8, f"Months to clear the {w}", months, months, "0")],
        pay_ref=lambda t: "=$C$7" if t <= k else "=$I$5", rate_ref=lambda t: "$C$6",
        c10=(f"Saving over the term of the {w} (£)", f"={orig:.2f}-SUM(C{FIRST_ROW}:C{last})", orig - total),
        answer="C10", grid_rows=n)
    return dict(text=text, answer=saving, steps=steps, worked=worked, meta=meta)


def holiday_case(p):
    P, i, n, R, F, orig, name, w = p["P"], p["i"], p["n"], p["R"], p["F"], p["orig"], p["name"], p["kind"]["word"]
    k = _change_month(n)
    h = random.choice([x for x in (1, 2, 3, 3, 6) if x <= n // 8])
    s = k + 1
    rows = run_schedule(P, lambda t: (i, 0 if s <= t < s + h else R))
    last_pay, months = round(rows[-1][1], 2), len(rows)
    total = sum(r[1] for r in rows)
    extra = round(total - orig, 2)
    held = round(rows[s + h - 2][4], 2)
    span = f"month {s}" if h == 1 else f"months {s}–{s + h - 1}"
    text = (f"After {k} repayments, {name} takes a {h}-month payment holiday: no repayment is made in {span}, but "
            f"interest is still added each month. Repayments of £{R:,.2f} then resume until the {w} is cleared.\n\n"
            f"Download the spreadsheet below. Complete the schedule to find the final repayment (I7) and the number "
            f"of months (I8), then determine how much extra the payment holiday costs {name} altogether (C10). Save "
            f"the file and upload it.")
    steps = [{"prompt": f"{w.capitalize()} outstanding at the end of month {k} (£)", "answer": round(rows[k - 1][4], 2)},
             {"prompt": f"{w.capitalize()} outstanding at the end of the holiday, month {s + h - 1} (£)", "answer": held},
             {"prompt": f"In which month is the {w} cleared?", "answer": months},
             {"prompt": "Final repayment (£)", "answer": last_pay},
             {"prompt": "Extra cost of the payment holiday (£)", "answer": extra}]
    worked = [f"{span.capitalize()}: repayment = 0, so the capital content is negative and the outstanding grows by "
              f"that month's interest",
              f"Outstanding: £{rows[k - 1][4]:,.2f} at the end of month {k}, £{held:,.2f} at the end of the holiday",
              f"Repayments of £{R:,.2f} resume; the {w} is cleared in month {months} with a last repayment of "
              f"£{last_pay:,.2f}",
              f"Total repaid now = {months - h - 1} × {R:,.2f} + {last_pay:,.2f} = £{total:,.2f}",
              f"Total repaid originally = {n - 1} × {R:,.2f} + {F:,.2f} = £{orig:,.2f}",
              f"Extra cost = £{total:,.2f} − £{orig:,.2f} = £{extra:,.2f}"]
    last = FIRST_ROW + months - 1
    meta = change_sheet(
        p, "payment holiday", rows,
        given={"C7": R, "C8": F, "C9": n},
        extras=[(4, "Payment holiday starts in month", s, "0"), (5, "Length of payment holiday (months)", h, "0")],
        fills=[(7, "Final repayment on the new plan (£)", f"=C{last}", last_pay, MONEY),
               (8, f"Months to clear the {w}", months, months, "0")],
        pay_ref=lambda t: 0 if s <= t < s + h else "=$C$7", rate_ref=lambda t: "$C$6",
        c10=("Extra cost of the payment holiday (£)", f"=SUM(C{FIRST_ROW}:C{last})-{orig:.2f}", total - orig),
        answer="C10", grid_rows=n + h + 12)
    return dict(text=text, answer=extra, steps=steps, worked=worked, meta=meta)


def rate_change_case(p):
    P, i, n, R, name, w = p["P"], p["i"], p["n"], p["R"], p["name"], p["kind"]["word"]
    k = 12 * random.choice([y for y in (1, 2, 3) if n - 12 * y >= 12])
    a2 = round(p["a"] + random.uniform(0.6, 2.4), 1)
    i2 = monthly_rate(a2 / 100)
    B = run_schedule(P, lambda t: (i, R), n=n)[k - 1][4]
    R2 = goal_seek_repayment(B, i2, n - k)
    rows = run_schedule(P, lambda t: (i, R) if t <= k else (i2, R2), n=n)
    F2 = round(rows[-1][1], 2)
    fixed = "first year" if k == 12 else f"first {k // 12} years"
    text = (f"The rate of {p['a']}% is fixed for the {fixed} only. From month "
            f"{k + 1} the annual effective rate rises to {a2}%, and {name}'s level monthly repayment is recalculated "
            f"so that the {w} is still cleared by the end of the original {n // 12}-year term.\n\nDownload the "
            f"spreadsheet below. Complete the schedule to determine the new level monthly repayment (cell I7) and "
            f"the new final repayment (I8). Save the file and upload it.")
    steps = [{"prompt": "New monthly effective rate, as a percentage (3 d.p.)", "answer": round(i2 * 100, 3)},
             {"prompt": f"{w.capitalize()} outstanding at the end of the fixed period, month {k} (£)", "answer": round(B, 2)},
             {"prompt": "New level monthly repayment, to 2 d.p. (£)", "answer": R2},
             {"prompt": "New final repayment (£)", "answer": F2}]
    worked = [f"Months 1–{k}: interest =ROUND($C$6*F…,2), repayment =$C$7 (£{R:,.2f})",
              f"Outstanding at the end of month {k}: £{B:,.2f}",
              f"I6 = (1 + {a2 / 100:g})^(1/12) − 1 = {i2 * 100:.3f}%; from month {k + 1}: interest =ROUND($I$6*F…,2), "
              f"repayment =$I$7",
              f"Goal Seek: month-{n} outstanding = 0 by changing I7 → £{R2:,.2f} (to 2 d.p.)",
              f"New final repayment = what is left in month {n} = £{F2:,.2f}"]
    last = FIRST_ROW + n - 1
    meta = change_sheet(
        p, "change of interest rate", rows,
        given={"C7": R, "C9": n}, blank=("B8", "B10"),
        extras=[(4, "New rate applies from month", k + 1, "0"), (5, "New annual effective rate", a2 / 100, RATE)],
        fills=[(6, "New monthly effective rate", "=(1+I5)^(1/12)-1", i2, RATE),
               (7, "New level monthly repayment (£)", R2, R2, MONEY),
               (8, "New final repayment (£)", f"=C{last}", F2, MONEY)],
        pay_ref=lambda t: "=$C$7" if t <= k else "=$I$7",
        rate_ref=lambda t: "$C$6" if t <= k else "$I$6",
        answer="I7", grid_rows=n)
    return dict(text=text, answer=R2, steps=steps, worked=worked, meta=meta)


CHANGE_NOTES = """- **Increased repayment:** from the month of the change, the repayment cell points at the new repayment
  (**=$I$5**); fill down until the outstanding reaches £0. Saving = total repaid originally − total repaid now
- **Payment holiday:** the repayment is **0** in the holiday months, so the capital content is negative and the
  outstanding **grows** by that month's interest; repayments then resume. Extra cost = total repaid now − total
  repaid originally
- **Change of interest rate:** from the month the new rate starts, the interest uses the new monthly rate
  (**ROUND($I$6\\*F…,2)**) and the repayment uses the new repayment (**=$I$7**); Goal Seek the last outstanding
  to **0** by changing I7, then round I7 to 2 d.p.
"""
