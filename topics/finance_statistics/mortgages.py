"""Higher Finance — Mortgages.

Mirrors Mortgages_Worksheet.docx (Higher Apps/Worksheets/Finance). Replaces the old flat-rate version:
the Higher Apps course uses effective (compound) rates only. SQA convention (2024 Q9, 2026 Q8 marking
instructions): monthly rate (1 + annual)^(1/12) − 1, interest = ROUND(rate × previous outstanding, 2);
paying the lender's maximum shortens the term and the saving is the difference in the TOTAL repaid;
affordability = repayments no more than 28% of monthly taxable income (after pension contributions);
loan-to-value = mortgage ÷ property value.
"""
import random

from openpyxl.styles import Font

from core.engine.spreadsheet_solution import excel_round, solution_metadata, workbook_bytes
from core.models.question_model import Question, make_part, multipart_worked_solution
from topics.finance_statistics.loan_spreadsheet import (
    ANSWER_FILL, FIRST_ROW, MONEY, MORTGAGE_LABELS, RATE, clear_schedule, level_repayment, monthly_rate,
    schedule, schedule_sheet, target_repayment,
)

TOPIC = "Finance"
QTYPE = "Mortgages"
_NAMES = ["Morag", "Ceit", "Domhnall", "Janet", "Fraser", "Kirsty", "Calum", "Iona", "Ewan", "Niall"]

NOTES = """
**Mortgages**

- **Monthly rate** = (1 + annual rate)^(1/12) − 1 — never annual ÷ 12, and never type a rounded rate into a spreadsheet
- **Each month:** interest = rate × mortgage outstanding (to 2 d.p.); capital = repayment − interest;
  outstanding = previous outstanding − capital
- **Paying the lender's maximum** shortens the term; the **saving** = total repaid on the original plan −
  total repaid on the new plan (not the difference in the monthly repayments)
- **Target balance** at the end of a fixed rate: Goal Seek the outstanding at the end of the fixed period to the
  **target** (not to 0), then round the repayment **up**
- **Affordability:** repayments should be no more than **28% of monthly taxable income** (salary after pension
  contributions). Always give the figure as your reason.
- **Loan-to-value** = mortgage ÷ property value (not deposit ÷ value)

**Example:** Lenders recommend that mortgage repayments are no more than 28% of monthly taxable income. Iona earns
£34,800 a year and pays 5% into her pension. Her repayment would be £720. Is it affordable?

- Taxable income = 34,800 × 0.95 = £33,060 a year = £2,755.00 a month
- 28% of £2,755.00 = £771.40
- £720 < £771.40, so **yes — affordable**
"""


def _monthly(a):
    return (1 + a) ** (1 / 12) - 1


def _clear(P, i, R):
    bal, rows = round(P, 2), []
    while bal > 0 and len(rows) < 600:
        it = round(i * bal + 1e-9, 2)
        pay = min(R, round(bal + it, 2))
        bal = round(bal - (pay - it), 2)
        rows.append((pay, it, bal))
    return rows


def _level(P, i, n):
    R = round(P * i / (1 - (1 + i) ** -n), 2)
    bal = round(P, 2)
    for _ in range(n - 1):
        bal = round(bal - (R - round(i * bal + 1e-9, 2)), 2)
    F = round(bal + round(i * bal + 1e-9, 2), 2)
    return R, F


# (context, amount lo, hi, step, annual % lo, hi, years options)
_MORTGAGES = [("a house", 90000, 240000, 5000, 3.9, 5.9, [20, 25, 30]),
              ("an extension to their croft house", 15000, 35000, 1000, 3.0, 4.5, [4, 5, 6]),
              ("a flat in Stornoway", 70000, 140000, 5000, 4.1, 5.6, [20, 25])]


def generate_mortgages_l1():
    """Mortgage outstanding at the end of month 2, by hand."""
    name = random.choice(_NAMES); ctx, lo, hi, st, rlo, rhi, yrs = random.choice(_MORTGAGES)
    P = random.randrange(lo, hi + 1, st); a = round(random.uniform(rlo, rhi), 1); i = _monthly(a / 100)
    R = round(P * i / (1 - (1 + i) ** -(12 * random.choice(yrs))) / 5) * 5 + 5
    it1 = round(i * P + 1e-9, 2); b1 = round(P - (R - it1), 2)
    it2 = round(i * b1 + 1e-9, 2); b2 = round(b1 - (R - it2), 2)
    q = (f"{name} takes out a £{P:,} mortgage for {ctx}. The annual effective rate of interest is {a}%, and repayments of "
         f"£{R:,} are made at the end of each month.\n\nCalculate the mortgage outstanding at the end of month 2.")
    return Question(question_text=q, correct_answer=b2, topic=TOPIC, question_type=QTYPE,
                    scaffold_steps=[{"prompt": "Monthly effective rate, as a percentage (3 d.p.)", "answer": round(i * 100, 3)},
                                    {"prompt": "Mortgage outstanding at the end of month 1 (£)", "answer": b1},
                                    {"prompt": "Interest content of repayment 2 (£)", "answer": it2}],
                    worked_solution=[f"Monthly rate = {1 + a / 100:.3f}^(1/12) − 1 = {i * 100:.3f}%",
                                     f"Month 1: interest {it1:,.2f}, outstanding {b1:,.2f}",
                                     f"Month 2: interest {it2:,.2f}, outstanding £{b2:,.2f}"],
                    notes=NOTES)


def generate_mortgages_l2():
    """Loan-to-value ratio."""
    name = random.choice(_NAMES)
    V = random.randrange(120000, 320001, 5000)
    dep = random.choice([0.05, 0.10, 0.15, 0.20, 0.25, 0.30])
    D = round(V * dep)
    ltv = round((V - D) / V * 100, 2)
    q = (f"{name} buys a house valued at £{V:,}, paying a deposit of £{D:,} and taking a mortgage for the rest.\n\n"
         f"Calculate the loan-to-value ratio, as a percentage.")
    return Question(question_text=q, correct_answer=ltv, topic=TOPIC, question_type=QTYPE,
                    scaffold_steps=[{"prompt": "Mortgage amount (£)", "answer": V - D}],
                    worked_solution=[f"Mortgage = {V:,} − {D:,} = £{V - D:,}",
                                     f"Loan-to-value = {V - D:,} ÷ {V:,} = {ltv:g}%  (not deposit ÷ value)"],
                    notes=NOTES)


def generate_mortgages_l3():
    """Maximum affordable repayment under the 28% guideline (styled on 2026 Q8(c))."""
    name = random.choice(_NAMES)
    sal = random.randrange(22000, 52001, 100); pen = random.choice([0.03, 0.04, 0.05, 0.06, 0.08])
    tax_m = sal * (1 - pen) / 12; lim = round(0.28 * tax_m, 2)
    q = (f"Lenders recommend that mortgage repayments are no more than 28% of monthly taxable income (salary after "
         f"pension contributions). {name} earns £{sal:,} a year and pays {pen:.0%} of it into a pension.\n\n"
         f"Calculate the maximum monthly repayment the lender would consider affordable.")
    return Question(question_text=q, correct_answer=lim, topic=TOPIC, question_type=QTYPE,
                    scaffold_steps=[{"prompt": "Annual taxable income (£)", "answer": round(sal * (1 - pen), 2)},
                                    {"prompt": "Monthly taxable income (£)", "answer": round(tax_m, 2)}],
                    worked_solution=[f"Taxable income = {sal:,} × {1 - pen:.2f} = £{sal * (1 - pen):,.2f} a year",
                                     f"Monthly = £{tax_m:,.2f};  28% = £{lim:,.2f}"],
                    notes=NOTES)


def generate_mortgages_l4():
    """Saving from paying the lender's maximum repayment (styled on 2024 Q9(c))."""
    name = random.choice(_NAMES)
    P = random.randrange(15000, 40001, 1000); a = round(random.uniform(3.0, 4.8), 1); yrs = random.choice([4, 5, 6])
    i = _monthly(a / 100); R, F = _level(P, i, yrs * 12)
    mx = int(R // 50 + 2) * 50
    inc = _clear(P, i, mx)
    orig = round((yrs * 12 - 1) * R + F, 2); new = round(sum(r[0] for r in inc), 2)
    q = (f"{name}'s £{P:,} mortgage over {yrs} years has {yrs * 12 - 1} level repayments of £{R:,.2f} and a final repayment "
         f"of £{F:,.2f}. By paying the lender's maximum of £{mx} a month instead, the mortgage is cleared with "
         f"{len(inc) - 1} repayments of £{mx} and a final repayment of £{inc[-1][0]:,.2f}.\n\n"
         f"Calculate how much {name} saves over the term of the mortgage.")
    return Question(question_text=q, correct_answer=round(orig - new, 2), topic=TOPIC, question_type=QTYPE,
                    scaffold_steps=[{"prompt": "Total repaid on the original plan (£)", "answer": orig},
                                    {"prompt": "Total repaid paying the maximum (£)", "answer": new}],
                    worked_solution=[f"Original: {yrs * 12 - 1} × {R:,.2f} + {F:,.2f} = £{orig:,.2f}",
                                     f"Maximum: {len(inc) - 1} × {mx} + {inc[-1][0]:,.2f} = £{new:,.2f}",
                                     f"Saving = £{orig - new:,.2f}  (not {mx} − {R:,.2f})"],
                    notes=NOTES)


def generate_mortgages_l5():
    """Total interest on a mortgage plan."""
    name = random.choice(_NAMES); ctx, lo, hi, st, rlo, rhi, yrs = random.choice(_MORTGAGES)
    P = random.randrange(lo, hi + 1, st); a = round(random.uniform(rlo, rhi), 1); y = random.choice(yrs)
    R, F = _level(P, _monthly(a / 100), y * 12)
    total = round((y * 12 - 1) * R + F, 2)
    q = (f"{name}'s £{P:,} mortgage for {ctx} is repaid with {y * 12 - 1} level monthly repayments of £{R:,.2f} and a "
         f"final repayment of £{F:,.2f}.\n\nCalculate the total interest paid.")
    return Question(question_text=q, correct_answer=round(total - P, 2), topic=TOPIC, question_type=QTYPE,
                    scaffold_steps=[{"prompt": "Total repaid (£)", "answer": total}],
                    worked_solution=[f"Total repaid = {y * 12 - 1} × {R:,.2f} + {F:,.2f} = £{total:,.2f}",
                                     f"Total interest = {total:,.2f} − {P:,} = £{total - P:,.2f}"],
                    notes=NOTES)


# ---------------------------------------------------------------------------
# Spreadsheet questions (worksheet Sections 1–3), in the same layout as Mortgages_Worksheet.xlsx:
# C4 amount, C5 annual rate, C6 monthly rate, C7 repayment, C8 final repayment, C9 term,
# C10 target; row 11 headers, row 12 month 0, rows 13… months 1…n.
# ---------------------------------------------------------------------------

SPREADSHEET_NOTES = """
**Mortgage schedules in a spreadsheet**

Same layout as the worksheet's spreadsheet questions:

- Monthly rate: **C6 = (1+C5)^(1/12)−1** — keep the formula; never type a rounded rate
- **F12 = C4**, then month 1 and fill down: interest **D13 = ROUND($C$6\\*F12,2)**, capital **E13 = C13−D13**,
  outstanding **F13 = F12−E13**, repayment **C13 = $C$7**
- **Level repayment:** Goal Seek the last outstanding to **0** by changing C7, round C7 to 2 d.p.; the last
  month repays what is left: **F(last−1)+ROUND($C$6\\*F(last−1),2)**
- **Paying the maximum:** put the maximum in C7 and fill down until the outstanding reaches £0; the last
  repayment is only what is left. Saving = total repaid originally − total repaid now
- **Target balance:** schedule for the fixed-rate months only; Goal Seek the last outstanding to the **target**
  (not 0), then round the repayment **up** to the next penny
- **Increased repayment:** from the month of the change, the repayment cell points at the new repayment
  (**=$I$5**); fill down until the outstanding reaches £0. Saving = total repaid originally − total repaid now
- **Payment holiday:** the repayment is **0** in the holiday months, so the capital content is negative and the
  outstanding **grows** by that month's interest; repayments then resume. Extra cost = total repaid now − total
  repaid originally
- **Change of interest rate:** from the month the new rate starts, the interest uses the new monthly rate
  (**ROUND($I$6\\*F…,2)**) and the repayment uses the new repayment (**=$I$7**); Goal Seek the last outstanding
  to **0** by changing I7, then round I7 to 2 d.p.

Save the file before uploading.
"""


def _mortgage_sheet(title, mode, P, a, i, rows, R, final=None, target=None, n_orig=None, orig_total=None):
    """Question workbook bytes, graded cell and completed-solution metadata for one mortgage
    spreadsheet. `rows` is the solution schedule (month, repayment, interest, capital, outstanding)."""
    n = len(rows)
    grid_rows = n_orig or n                 # the question sheet gets the original term's month numbers
    last = FIRST_ROW + n - 1

    def base(solution):
        wb, ws = schedule_sheet(title, "Mortgage", grid_rows if not solution else n, MORTGAGE_LABELS)
        ws["C4"], ws["C5"] = P, a / 100
        if target is not None:
            ws["C10"] = target
        return wb, ws

    to_fill = ["C6", "F12"] + [f"{c}{r}" for r in range(FIRST_ROW, FIRST_ROW + grid_rows) for c in "CDEF"]
    if mode == "repayment":
        to_fill += ["C7", "C8"]
        answer = "C7"
    elif mode == "target":
        to_fill += ["C7"]
        answer = "C7"
    else:                                   # maximum
        to_fill += ["C8", "C9", "C10"]
        answer = "C10"

    wb, ws = base(False)
    if mode == "maximum":
        ws["C7"] = R
        ws["C9"] = None
        ws["B10"] = "Saving over the term of the mortgage (£)"
    for ref in to_fill:
        ws[ref].fill = ANSWER_FILL
    ws["A2"] = "Fill in the yellow cells, save the file, then upload it."
    ws["A2"].font = Font(italic=True)
    question_bytes = workbook_bytes(wb)

    wb, ws = base(True)
    ws["A2"] = "Completed solution: the yellow cells show the formulas used."
    ws["A2"].font = Font(italic=True)
    ws["C6"] = "=(1+C5)^(1/12)-1"
    ws["F12"] = "=C4"
    values = {"C6": i, "F12": P}
    ws["C7"] = R
    if mode == "repayment":
        ws["D7"] = "← Goal Seek: last 'Mortgage outstanding' = 0 by changing C7, then round"
    elif mode == "target":
        ws["D7"] = "← Goal Seek: month-n 'Mortgage outstanding' = target by changing C7, then round UP"
    for t, pay, interest, capital, bal in rows:
        r = FIRST_ROW + t - 1
        clears = mode != "target" and t == n
        ws[f"C{r}"] = f"=F{r - 1}+ROUND($C$6*F{r - 1},2)" if clears else "=$C$7"
        ws[f"D{r}"] = f"=ROUND($C$6*F{r - 1},2)"
        ws[f"E{r}"] = f"=C{r}-D{r}"
        ws[f"F{r}"] = f"=F{r - 1}-E{r}"
        values.update({f"C{r}": pay, f"D{r}": interest, f"E{r}": capital, f"F{r}": bal})
    if mode != "target":
        ws["C8"] = f"=C{last}"
        values["C8"] = rows[-1][1]
    if mode == "maximum":
        ws["C9"] = n
        ws["D8"] = f"← the schedule reaches £0 in month {n}; this is what is left"
        ws["B10"] = "Saving over the term of the mortgage (£)"
        ws["C10"] = f"={orig_total:.2f}-SUM(C{FIRST_ROW}:C{last})"
        values["C10"] = orig_total - sum(r[1] for r in rows)
        ws["D10"] = f"← total repaid on the original plan (£{orig_total:,.2f}) − total repaid now"
    filled = [ref for ref in to_fill if int(ref[1:]) <= last]
    for ref in filled:
        ws[ref].fill = ANSWER_FILL
    meta = solution_metadata(wb, ws, values, filled, min_row=4, max_row=last, min_col=2, max_col=6,
                             filename="mortgage_schedule_solution.xlsx")
    return {"spreadsheet_bytes": question_bytes, "spreadsheet_filename": "mortgage_schedule.xlsx",
            "spreadsheet_answer_cell": ("Mortgage", answer), **meta}


_SHEET_MORTGAGES = [("is extending the croft house", 15000, 35000, 500, 3.0, 4.8, [4, 5, 6]),
                    ("is building a workshop and polytunnel on the croft", 20000, 40000, 500, 3.2, 4.6, [5, 6]),
                    ("is renovating the roof of the house", 12000, 25000, 500, 3.5, 5.0, [3, 4, 5])]


def generate_mortgages_l6():
    """Spreadsheet: Goal Seek the level repayment and final repayment (worksheet Section 1)."""
    name = random.choice(_NAMES); ctx, lo, hi, st, rlo, rhi, yrs = random.choice(_SHEET_MORTGAGES)
    P = random.randrange(lo, hi + 1, st); a = round(random.uniform(rlo, rhi), 1); n = 12 * random.choice(yrs)
    i = _monthly(a / 100)
    R = round(level_repayment(P, i, n), 2)
    rows = schedule(P, i, R, n)
    F = round(rows[-1][1], 2)
    q = (f"{name} {ctx}. {name} is offered a £{P:,} mortgage with an annual effective rate of interest of {a}% over "
         f"{n // 12} years, with level repayments at the end of each month.\n\nDownload the spreadsheet below. Complete "
         f"the mortgage schedule to determine the level monthly repayment (cell C7) and the final repayment (cell C8), "
         f"then save it and upload it here.")
    return Question(
        question_text=q, correct_answer=R, topic=TOPIC, question_type=QTYPE,
        scaffold_steps=[{"prompt": "Monthly effective rate, as a percentage (3 d.p.)", "answer": round(i * 100, 3)},
                        {"prompt": "Interest content of the first repayment (£)", "answer": rows[0][2]},
                        {"prompt": "Level monthly repayment, to 2 d.p. (£)", "answer": R},
                        {"prompt": "Final repayment (£)", "answer": F}],
        worked_solution=[f"C6 = (1 + {a / 100:g})^(1/12) − 1 = {i * 100:.3f}% (keep the formula — don't type the rounded rate)",
                         f"D13 = ROUND($C$6*F12,2) = {rows[0][2]:,.2f}; fill down to month {n}",
                         f"Goal Seek: month-{n} outstanding = 0 by changing C7 → £{R:,.2f} for months 1–{n - 1}",
                         f"Final repayment = what is left in month {n} = £{F:,.2f}"],
        notes=SPREADSHEET_NOTES,
        metadata=_mortgage_sheet(f"{name}'s Mortgage", "repayment", P, a, i, rows, R))


def generate_mortgages_l7():
    """Spreadsheet: pay the lender's maximum — final repayment and saving (worksheet Section 2)."""
    name = random.choice(_NAMES); ctx, lo, hi, st, rlo, rhi, yrs = random.choice(_SHEET_MORTGAGES)
    P = random.randrange(lo, hi + 1, st); a = round(random.uniform(rlo, rhi), 1); n = 12 * random.choice(yrs)
    i = _monthly(a / 100)
    R = round(level_repayment(P, i, n), 2)
    F = round(schedule(P, i, R, n)[-1][1], 2)
    orig = round((n - 1) * R + F, 2)
    mx = int(R // 50 + 2) * 50
    rows = clear_schedule(P, i, mx)
    last_pay = round(rows[-1][1], 2)
    saving = round(orig - sum(r[1] for r in rows), 2)
    q = (f"{name}'s £{P:,} mortgage over {n // 12} years has an annual effective rate of {a}%. The level repayment is "
         f"£{R:,.2f} with a final repayment of £{F:,.2f}. The lender's maximum monthly repayment is £{mx}, and {name} "
         f"chooses to pay this to reduce the term.\n\nDownload the spreadsheet below. Complete the schedule for the "
         f"reduced term to find the final repayment (C8) and the number of months (C9), then determine how much money "
         f"this saves over the term of the mortgage (C10). Save the file and upload it here.")
    return Question(
        question_text=q, correct_answer=saving, topic=TOPIC, question_type=QTYPE,
        scaffold_steps=[{"prompt": "Total repaid on the original plan (£)", "answer": orig},
                        {"prompt": f"In which month does the mortgage reach £0 paying £{mx}?", "answer": len(rows)},
                        {"prompt": "Final repayment (£)", "answer": last_pay},
                        {"prompt": "Saving over the term (£)", "answer": saving}],
        worked_solution=[f"Put £{mx} in C7 and fill down: the mortgage is cleared in month {len(rows)}; the last "
                         f"repayment is only what is left: £{last_pay:,.2f}",
                         f"Total repaid originally = {n - 1} × {R:,.2f} + {F:,.2f} = £{orig:,.2f}",
                         f"Total repaid now = {len(rows) - 1} × {mx} + {last_pay:,.2f} = £{orig - saving:,.2f}",
                         f"Saving = £{saving:,.2f}  (not {mx} − {R:,.2f})"],
        notes=SPREADSHEET_NOTES,
        metadata=_mortgage_sheet(f"{name}'s Mortgage — paying the maximum", "maximum", P, a, i, rows, mx,
                                 n_orig=n, orig_total=orig))


def generate_mortgages_l8():
    """Spreadsheet: minimum repayment to reach a target balance by the end of a fixed rate (Section 3)."""
    name = random.choice(_NAMES)
    P = random.randrange(80000, 220001, 1000); a = round(random.uniform(3.9, 5.6), 1); y = random.choice([2, 3, 5])
    n = 12 * y
    target = int(P * random.uniform(0.86, 0.96) / 500) * 500
    i = _monthly(a / 100)
    R = target_repayment(P, i, n, target)
    rows = schedule(P, i, R, n, R)
    end = rows[-1][4]
    q = (f"{name} is applying for a £{P:,} mortgage. The fixed annual effective rate is {a}% for the first {y} years, with "
         f"level repayments at the end of every month. {name} wants the mortgage outstanding to be no more than "
         f"£{target:,} at the end of the fixed-rate period.\n\nDownload the spreadsheet below. Complete the schedule to "
         f"determine the minimum level monthly repayment (cell C7), then save it and upload it here.")
    return Question(
        question_text=q, correct_answer=R, topic=TOPIC, question_type=QTYPE,
        scaffold_steps=[{"prompt": "Monthly effective rate, as a percentage (3 d.p.)", "answer": round(i * 100, 3)},
                        {"prompt": "How many months does the schedule need?", "answer": n},
                        {"prompt": "Minimum level monthly repayment, rounded up to the penny (£)", "answer": R}],
        worked_solution=[f"Build the schedule for {n} months only (the fixed-rate period)",
                         f"Goal Seek: month-{n} outstanding = {target:,} (NOT 0) by changing C7",
                         f"Round the repayment UP to the next penny: £{R:,.2f} → month-{n} outstanding £{end:,.2f}"],
        notes=SPREADSHEET_NOTES,
        metadata=_mortgage_sheet(f"{name}'s Fixed-Rate Mortgage", "target", P, a, i, rows, R, target=target))


# ---------------------------------------------------------------------------
# Spreadsheet questions on a change to the plan part-way through: an increased repayment, a payment
# holiday, a change of interest rate. Same layout, with the change's details in H4:I8.
# ---------------------------------------------------------------------------

def _run(P, terms, n=None, max_months=600):
    """Schedule rows (month, repayment, interest, capital, outstanding) where terms(t) gives month t's
    (monthly rate, repayment). Without n, runs until cleared and the last repayment is only what is
    left; with n, month n repays whatever is left (the level-repayment final row). Same float-faithful
    arithmetic as loan_spreadsheet.schedule()."""
    bal, rows = P, []
    for t in range(1, max_months + 1):
        i, R = terms(t)
        interest = excel_round(i * bal)
        pay = bal + interest if (t == n or (n is None and R >= bal + interest)) else R
        bal = bal - (pay - interest)
        rows.append((t, pay, interest, pay - interest, bal))
        if t == n or (n is None and bal <= 1e-9):
            break
    return rows


def _goal_seek(B, i, m):
    """The repayment (to 2 d.p.) that Goal Seek finds to clear balance B in m months at monthly rate i."""
    def end(R):
        bal = B
        for _ in range(m):
            bal -= R - excel_round(i * bal)
        return bal
    lo, hi = 0.0, B * 2
    for _ in range(80):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if end(mid) > 0 else (lo, mid)
    return excel_round((lo + hi) / 2)


def _change_sheet(title, P, a, rows, *, given, extras, fills, pay_ref, rate_ref, answer, grid_rows,
                  c10=None, blank=()):
    """Question workbook bytes, graded cell and completed-solution metadata for a plan that changes
    part-way. `given` — C-column values shown to pupils; `extras` — (row, label, value, format) given in
    H/I; `fills` — (row, label, formula_or_value, value, format) pupils complete in H/I; pay_ref(t) /
    rate_ref(t) — month t's repayment formula and rate cell; `c10` — (label, formula, value)."""
    n = len(rows)
    last = FIRST_ROW + n - 1

    def base(months):
        wb, ws = schedule_sheet(title, "Mortgage", months, MORTGAGE_LABELS)
        ws.column_dimensions["H"].width = 36
        ws.column_dimensions["I"].width = 16
        ws["C4"], ws["C5"] = P, a / 100
        for ref, value in given.items():
            ws[ref] = value
        for ref in blank:
            ws[ref] = None
        for row, label, value, fmt in extras:
            ws[f"H{row}"], ws[f"I{row}"] = label, value
            ws[f"I{row}"].number_format = fmt
        for row, label, _, _, fmt in fills:
            ws[f"H{row}"] = label
            ws[f"I{row}"].number_format = fmt
        if c10:
            ws["B10"] = c10[0]
        return wb, ws

    to_fill = ["C6", "F12"] + [f"I{row}" for row, *_ in fills] + (["C10"] if c10 else [])
    wb, ws = base(grid_rows)
    for ref in to_fill + [f"{c}{r}" for r in range(FIRST_ROW, FIRST_ROW + grid_rows) for c in "CDEF"]:
        ws[ref].fill = ANSWER_FILL
    ws["A2"] = "Fill in the yellow cells, save the file, then upload it."
    ws["A2"].font = Font(italic=True)
    question_bytes = workbook_bytes(wb)

    wb, ws = base(n)
    ws["A2"] = "Completed solution: the yellow cells show the formulas used."
    ws["A2"].font = Font(italic=True)
    ws["C6"], ws["F12"] = "=(1+C5)^(1/12)-1", "=C4"
    values = {"C6": monthly_rate(a / 100), "F12": P}
    for row, _, content, value, _ in fills:
        ws[f"I{row}"] = content
        values[f"I{row}"] = value
    for t, pay, interest, capital, bal in rows:
        r = FIRST_ROW + t - 1
        rate = rate_ref(t)
        ws[f"C{r}"] = f"=F{r - 1}+ROUND({rate}*F{r - 1},2)" if t == n else pay_ref(t)
        ws[f"D{r}"] = f"=ROUND({rate}*F{r - 1},2)"
        ws[f"E{r}"] = f"=C{r}-D{r}"
        ws[f"F{r}"] = f"=F{r - 1}-E{r}"
        values.update({f"C{r}": pay, f"D{r}": interest, f"E{r}": capital, f"F{r}": bal})
    if c10:
        ws["C10"] = c10[1]
        values["C10"] = c10[2]
    filled = to_fill + [f"{c}{r}" for r in range(FIRST_ROW, last + 1) for c in "CDEF"]
    for ref in filled:
        ws[ref].fill = ANSWER_FILL
    meta = solution_metadata(wb, ws, values, filled, min_row=4, max_row=last, min_col=2, max_col=9,
                             filename="mortgage_schedule_solution.xlsx")
    return {"spreadsheet_bytes": question_bytes, "spreadsheet_filename": "mortgage_schedule.xlsx",
            "spreadsheet_answer_cell": ("Mortgage", answer), **meta}


def _sheet_plan():
    """A random mortgage from _SHEET_MORTGAGES with its level repayment R and final repayment F."""
    name = random.choice(_NAMES); ctx, lo, hi, st, rlo, rhi, yrs = random.choice(_SHEET_MORTGAGES)
    P = random.randrange(lo, hi + 1, st); a = round(random.uniform(rlo, rhi), 1); n = 12 * random.choice(yrs)
    i = _monthly(a / 100)
    R = round(level_repayment(P, i, n), 2)
    F = round(schedule(P, i, R, n)[-1][1], 2)
    return dict(name=name, ctx=ctx, P=P, a=a, i=i, n=n, R=R, F=F, orig=round((n - 1) * R + F, 2))


def _plan_intro(p):
    return (f"{p['name']} {p['ctx']}. {p['name']}'s £{p['P']:,} mortgage over {p['n'] // 12} years has an annual "
            f"effective rate of {p['a']}%, with level monthly repayments of £{p['R']:,.2f} and a final repayment of "
            f"£{p['F']:,.2f}.")


def _increase_case(p):
    P, i, n, R, F, orig, name = p["P"], p["i"], p["n"], p["R"], p["F"], p["orig"], p["name"]
    k = random.choice([m for m in (12, 18, 24, 30, 36) if m <= n - 24])
    inc = random.choice([25, 50, 75, 100, 150])
    R2 = round(R + inc, 2)
    rows = _run(P, lambda t: (i, R if t <= k else R2))
    last_pay, months = round(rows[-1][1], 2), len(rows)
    total = sum(r[1] for r in rows)
    saving = round(orig - total, 2)
    text = (f"After {k} repayments, {name} increases the monthly repayment by £{inc} to £{R2:,.2f}, starting with "
            f"repayment {k + 1}, and keeps paying this until the mortgage is cleared.\n\nDownload the spreadsheet "
            f"below. Complete the schedule to find the final repayment (I7) and the number of months (I8), then "
            f"determine how much money this saves over the term of the mortgage (C10). Save the file and upload it.")
    steps = [{"prompt": "Total repaid on the original plan (£)", "answer": orig},
             {"prompt": f"Mortgage outstanding at the end of month {k} (£)", "answer": round(rows[k - 1][4], 2)},
             {"prompt": f"In which month does the mortgage reach £0 paying £{R2:,.2f}?", "answer": months},
             {"prompt": "Final repayment (£)", "answer": last_pay},
             {"prompt": "Saving over the term (£)", "answer": saving}]
    worked = [f"Months 1–{k}: repayment =$C$7 (£{R:,.2f}); from month {k + 1}: =$I$5 (£{R2:,.2f})",
              f"Outstanding at the end of month {k}: £{rows[k - 1][4]:,.2f}",
              f"Fill down: the mortgage is cleared in month {months}; the last repayment is only what is left: "
              f"£{last_pay:,.2f}",
              f"Total repaid originally = {n - 1} × {R:,.2f} + {F:,.2f} = £{orig:,.2f}",
              f"Total repaid now = {k} × {R:,.2f} + {months - k - 1} × {R2:,.2f} + {last_pay:,.2f} = £{total:,.2f}",
              f"Saving = £{saving:,.2f}"]
    last = FIRST_ROW + months - 1
    meta = _change_sheet(
        f"{name}'s Mortgage — increased repayment", p["P"], p["a"], rows,
        given={"C7": R, "C8": F, "C9": n},
        extras=[(4, "Repayment increases from month", k + 1, "0"), (5, "New monthly repayment (£)", R2, MONEY)],
        fills=[(7, "Final repayment on the new plan (£)", f"=C{last}", last_pay, MONEY),
               (8, "Months to clear the mortgage", months, months, "0")],
        pay_ref=lambda t: "=$C$7" if t <= k else "=$I$5", rate_ref=lambda t: "$C$6",
        c10=("Saving over the term of the mortgage (£)", f"={orig:.2f}-SUM(C{FIRST_ROW}:C{last})", orig - total),
        answer="C10", grid_rows=n)
    return dict(text=text, answer=saving, steps=steps, worked=worked, meta=meta)


def _holiday_case(p):
    P, i, n, R, F, orig, name = p["P"], p["i"], p["n"], p["R"], p["F"], p["orig"], p["name"]
    k = random.choice([m for m in (12, 18, 24, 30, 36) if m <= n - 18])
    h = random.choice([2, 3, 3, 6])
    s = k + 1
    rows = _run(P, lambda t: (i, 0 if s <= t < s + h else R))
    last_pay, months = round(rows[-1][1], 2), len(rows)
    total = sum(r[1] for r in rows)
    extra = round(total - orig, 2)
    held = round(rows[s + h - 2][4], 2)
    text = (f"After {k} repayments, {name} takes a {h}-month payment holiday: no repayments are made in months "
            f"{s}–{s + h - 1}, but interest is still added each month. Repayments of £{R:,.2f} then resume until "
            f"the mortgage is cleared.\n\nDownload the spreadsheet below. Complete the schedule to find the final "
            f"repayment (I7) and the number of months (I8), then determine how much extra the payment holiday "
            f"costs {name} altogether (C10). Save the file and upload it.")
    steps = [{"prompt": f"Mortgage outstanding at the end of month {k} (£)", "answer": round(rows[k - 1][4], 2)},
             {"prompt": f"Mortgage outstanding at the end of the holiday, month {s + h - 1} (£)", "answer": held},
             {"prompt": "In which month is the mortgage cleared?", "answer": months},
             {"prompt": "Final repayment (£)", "answer": last_pay},
             {"prompt": "Extra cost of the payment holiday (£)", "answer": extra}]
    worked = [f"Months {s}–{s + h - 1}: repayment = 0, so the capital content is negative and the outstanding "
              f"grows by each month's interest",
              f"Outstanding: £{rows[k - 1][4]:,.2f} at the end of month {k}, £{held:,.2f} at the end of the holiday",
              f"Repayments of £{R:,.2f} resume; the mortgage is cleared in month {months} with a last repayment of "
              f"£{last_pay:,.2f}",
              f"Total repaid now = {months - h - 1} × {R:,.2f} + {last_pay:,.2f} = £{total:,.2f}",
              f"Total repaid originally = {n - 1} × {R:,.2f} + {F:,.2f} = £{orig:,.2f}",
              f"Extra cost = £{total:,.2f} − £{orig:,.2f} = £{extra:,.2f}"]
    last = FIRST_ROW + months - 1
    meta = _change_sheet(
        f"{name}'s Mortgage — payment holiday", p["P"], p["a"], rows,
        given={"C7": R, "C8": F, "C9": n},
        extras=[(4, "Payment holiday starts in month", s, "0"), (5, "Length of payment holiday (months)", h, "0")],
        fills=[(7, "Final repayment on the new plan (£)", f"=C{last}", last_pay, MONEY),
               (8, "Months to clear the mortgage", months, months, "0")],
        pay_ref=lambda t: 0 if s <= t < s + h else "=$C$7", rate_ref=lambda t: "$C$6",
        c10=("Extra cost of the payment holiday (£)", f"=SUM(C{FIRST_ROW}:C{last})-{orig:.2f}", total - orig),
        answer="C10", grid_rows=n + h + 12)
    return dict(text=text, answer=extra, steps=steps, worked=worked, meta=meta)


def _rate_change_case(p):
    P, i, n, R, name = p["P"], p["i"], p["n"], p["R"], p["name"]
    k = 12 * random.choice([y for y in (1, 2, 3) if n - 12 * y >= 24])
    a2 = round(p["a"] + random.uniform(0.6, 2.4), 1)
    i2 = _monthly(a2 / 100)
    B = _run(P, lambda t: (i, R), n=n)[k - 1][4]
    R2 = _goal_seek(B, i2, n - k)
    rows = _run(P, lambda t: (i, R) if t <= k else (i2, R2), n=n)
    F2 = round(rows[-1][1], 2)
    text = (f"The rate of {p['a']}% is fixed for the first {k // 12} year{'s' if k > 12 else ''} only. From month "
            f"{k + 1} the annual effective rate rises to {a2}%, and {name}'s level monthly repayment is recalculated "
            f"so that the mortgage is still cleared by the end of the original {n // 12}-year term.\n\nDownload the "
            f"spreadsheet below. Complete the schedule to determine the new level monthly repayment (cell I7) and "
            f"the new final repayment (I8). Save the file and upload it.")
    steps = [{"prompt": "New monthly effective rate, as a percentage (3 d.p.)", "answer": round(i2 * 100, 3)},
             {"prompt": f"Mortgage outstanding at the end of the fixed period, month {k} (£)", "answer": round(B, 2)},
             {"prompt": "New level monthly repayment, to 2 d.p. (£)", "answer": R2},
             {"prompt": "New final repayment (£)", "answer": F2}]
    worked = [f"Months 1–{k}: interest =ROUND($C$6*F…,2), repayment =$C$7 (£{R:,.2f})",
              f"Outstanding at the end of month {k}: £{B:,.2f}",
              f"I6 = (1 + {a2 / 100:g})^(1/12) − 1 = {i2 * 100:.3f}%; from month {k + 1}: interest =ROUND($I$6*F…,2), "
              f"repayment =$I$7",
              f"Goal Seek: month-{n} outstanding = 0 by changing I7 → £{R2:,.2f} (to 2 d.p.)",
              f"New final repayment = what is left in month {n} = £{F2:,.2f}"]
    last = FIRST_ROW + n - 1
    meta = _change_sheet(
        f"{name}'s Mortgage — change of interest rate", p["P"], p["a"], rows,
        given={"C7": R, "C9": n}, blank=("B8", "B10"),
        extras=[(4, "New rate applies from month", k + 1, "0"), (5, "New annual effective rate", a2 / 100, RATE)],
        fills=[(6, "New monthly effective rate", "=(1+I5)^(1/12)-1", i2, RATE),
               (7, "New level monthly repayment (£)", R2, R2, MONEY),
               (8, "New final repayment (£)", f"=C{last}", F2, MONEY)],
        pay_ref=lambda t: "=$C$7" if t <= k else "=$I$7",
        rate_ref=lambda t: "$C$6" if t <= k else "$I$6",
        answer="I7", grid_rows=n)
    return dict(text=text, answer=R2, steps=steps, worked=worked, meta=meta)


def _rate_change_intro(p):
    return (f"{p['name']} {p['ctx']}. {p['name']} takes out a £{p['P']:,} mortgage over {p['n'] // 12} years at a "
            f"fixed annual effective rate of {p['a']}%, with level monthly repayments of £{p['R']:,.2f}.")


def _single(case, intro):
    p = _sheet_plan()
    c = case(p)
    return Question(question_text=f"{intro(p)} {c['text']}", correct_answer=c["answer"], topic=TOPIC,
                    question_type=QTYPE, scaffold_steps=c["steps"], worked_solution=c["worked"],
                    notes=SPREADSHEET_NOTES, metadata=c["meta"])


def generate_mortgages_l9():
    """Spreadsheet: increase the repayment part-way through — saving over the term."""
    return _single(_increase_case, _plan_intro)


def generate_mortgages_l10():
    """Spreadsheet: a payment holiday part-way through — extra cost."""
    return _single(_holiday_case, _plan_intro)


def generate_mortgages_l11():
    """Spreadsheet: the rate rises after a fixed period — Goal Seek the new level repayment."""
    return _single(_rate_change_case, _rate_change_intro)


def generate_mortgages_l12():
    """Exam style, in parts as 2024 Q9 / 2026 Q8: (a) spreadsheet level repayment, (b) affordability,
    (c) a spreadsheet change to the plan — an increased repayment, a payment holiday or a rate rise."""
    p = _sheet_plan()
    name, P, a, i, n, R, F = p["name"], p["P"], p["a"], p["i"], p["n"], p["R"], p["F"]
    case, intro = random.choice([(_increase_case, _plan_intro), (_holiday_case, _plan_intro),
                                 (_rate_change_case, _rate_change_intro)])
    rate_note = " (fixed for the first part of the term — see part (c))" if case is _rate_change_case else ""
    context = (f"{name} {p['ctx']} and is offered a £{P:,} mortgage over {n // 12} years at an annual effective "
               f"rate of {a}%{rate_note}, with level repayments at the end of each month.")

    rows = schedule(P, i, R, n)
    part_a = make_part(
        "(a)", "Download the spreadsheet below. Complete the mortgage schedule to determine the level monthly "
               "repayment (cell C7) and the final repayment (cell C8), then save it and upload it here.", R,
        scaffold_steps=[{"prompt": "Monthly effective rate, as a percentage (3 d.p.)", "answer": round(i * 100, 3)},
                        {"prompt": "Interest content of the first repayment (£)", "answer": rows[0][2]},
                        {"prompt": "Level monthly repayment, to 2 d.p. (£)", "answer": R}],
        worked_solution=[f"C6 = (1 + {a / 100:g})^(1/12) − 1 = {i * 100:.3f}%",
                         f"Goal Seek: month-{n} outstanding = 0 by changing C7 → £{R:,.2f}",
                         f"Final repayment = what is left in month {n} = £{F:,.2f}"])
    part_a.metadata.update(_mortgage_sheet(f"{name}'s Mortgage", "repayment", P, a, i, rows, R))

    pen = random.choice([0.03, 0.04, 0.05, 0.06, 0.08])
    while True:
        sal = round(R * random.uniform(0.8, 1.25) / 0.28 * 12 / (1 - pen) / 100) * 100
        tax_m = sal * (1 - pen) / 12; lim = round(0.28 * tax_m, 2)
        if abs(lim - R) > 5:
            break
    ok = R <= lim
    yes, no = "Yes — the repayment is affordable", "No — the repayment is not affordable"
    part_b = make_part(
        "(b)", f"Lenders recommend that mortgage repayments are no more than 28% of monthly taxable income (salary "
               f"after pension contributions). {name} earns £{sal:,} a year and pays {pen:.0%} of it into a pension. "
               f"Using your answer to (a), determine whether the lender would consider the mortgage affordable.",
        yes if ok else no, options=[yes, no],
        scaffold_steps=[{"prompt": "Monthly taxable income (£)", "answer": round(tax_m, 2)},
                        {"prompt": "28% of monthly taxable income (£)", "answer": lim}],
        worked_solution=[f"Taxable income = {sal:,} × {1 - pen:.2f} = £{sal * (1 - pen):,.2f} a year = "
                         f"£{tax_m:,.2f} a month",
                         f"28% of £{tax_m:,.2f} = £{lim:,.2f}",
                         f"£{R:,.2f} {'≤' if ok else '>'} £{lim:,.2f}, so {'yes — affordable' if ok else 'no — not affordable'}"])

    c = case(p)
    part_c = make_part("(c)", c["text"], c["answer"], scaffold_steps=c["steps"], worked_solution=c["worked"])
    part_c.metadata.update(c["meta"])

    parts = [part_a, part_b, part_c]
    return Question(question_text=context, correct_answer=c["answer"], topic=TOPIC, question_type=QTYPE,
                    parts=parts, worked_solution=multipart_worked_solution(parts), notes=SPREADSHEET_NOTES)


# The non-spreadsheet question types, grouped as one "Calculator Questions" entry in the app and
# weighted toward what the 2023–2026 past papers actually ask (see the comments).
_CALCULATOR_MIX = [
    (generate_mortgages_l1, 3),             # outstanding after 2 months by hand — as the 2024 Q1 loan
    (generate_mortgages_l3, 3),             # affordability, 28% of taxable income — 2026 Q8(c)
    (generate_mortgages_l4, 2),             # saving from paying the maximum — 2024 Q9(c)
    (generate_mortgages_l5, 2),             # total interest — as 2023 Q11(b)
    (generate_mortgages_l2, 1),             # loan-to-value — not yet examined
]


def generate_mortgages_calculator():
    generator = random.choices([g for g, _ in _CALCULATOR_MIX], weights=[w for _, w in _CALCULATOR_MIX])[0]
    return generator()


def generate_mortgages_question():
    return random.choice([generate_mortgages_calculator, generate_mortgages_l6, generate_mortgages_l7, generate_mortgages_l8,
                          generate_mortgages_l9, generate_mortgages_l10, generate_mortgages_l11,
                          generate_mortgages_l12])()
