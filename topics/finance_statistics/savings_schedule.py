import calendar
import io
import random
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models.question_model import Question

NOTES = """
**Savings Schedule (spreadsheet):**

A savings account is opened with an initial deposit, then grows by a fixed
**monthly payment** plus **interest** applied once a month, right before each
payment is made. The interest rate itself can **change partway through** —
an initial rate quoted **per month**, switching on a later date to a rate
quoted **per year**.

1. Convert any annual rate to its **monthly-equivalent** rate:
   monthly rate = (1 + annual rate)^(1/12) − 1
2. Each month: balance before payment = previous balance after
   + ROUND(rate × previous balance after, 2) — using whichever rate applies
   on that date
3. Balance after payment = balance before payment + the monthly payment
4. Carry the balance forward, month by month, until the date asked for

Download the spreadsheet below, fill in the yellow cells (working month by
month down the table), save it, then upload it here to check your answer.
"""

_NAMES = [
    "Anna", "Iain", "Ceitidh", "Fraser", "Niamh", "Callum",
    "Eilidh", "Grant", "Morven", "Ruaridh", "Skye", "Lachlan",
]

_PRONOUNS = {
    "Anna": ("she", "her"), "Iain": ("he", "his"), "Ceitidh": ("she", "her"),
    "Fraser": ("he", "his"), "Niamh": ("she", "her"), "Callum": ("he", "his"),
    "Eilidh": ("she", "her"), "Grant": ("he", "his"), "Morven": ("she", "her"),
    "Ruaridh": ("he", "his"), "Skye": ("they", "their"), "Lachlan": ("he", "his"),
}

_SCENARIOS = [
    "a new laptop and coding course fees",
    "new solar panels for the croft",
    "university accommodation costs",
    "a second-hand car",
    "a wedding",
    "a deposit on a first flat",
    "a gap year trip",
    "new kitchen appliances",
]

_DEPOSITS = list(range(200, 851, 50))
_PAYMENTS = list(range(100, 251, 10))
_INITIAL_MONTHLY_RATES = [0.14, 0.16, 0.18, 0.19, 0.21, 0.23, 0.25]
_ANNUAL_RATES = [2.2, 2.6, 2.9, 3.1, 3.4, 3.8]

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_ANSWER_FILL = PatternFill("solid", fgColor="FFFF00")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _ym_to_mi(year, month):
    return year * 12 + (month - 1)


def _mi_to_ym(mi):
    year, m0 = divmod(mi, 12)
    return year, m0 + 1


def _date_str(d):
    return f"{d.day} {calendar.month_name[d.month]} {d.year}"


def _fmt_money(x):
    return f"{x:,.2f}"


def _add_months(start_mi, n):
    year, month = _mi_to_ym(start_mi + n)
    return date(year, month, 1)


def _excel_round(x, dp=2):
    """Round half away from zero, matching Excel's ROUND() — Python's round()
    uses banker's rounding, which can disagree with what a student's actual
    spreadsheet computes."""
    quant = Decimal(1).scaleb(-dp)
    return float(Decimal(str(x)).quantize(quant, rounding=ROUND_HALF_UP))


def _last_day_of_month(d):
    last_day = calendar.monthrange(d.year, d.month)[1]
    return date(d.year, d.month, last_day)


def _build_workbook(scenario):
    wb = Workbook()
    ws = wb.active
    ws.title = "Savings"

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 22

    ws["A1"] = "Name:"
    ws["A2"] = "Class:"

    ws["A4"] = f"{scenario['name']}'s Savings Schedule"
    ws["A4"].font = Font(bold=True, size=13)

    ws["A5"] = "Enter your answers in the yellow cells. Other values are given."
    ws["A5"].font = Font(italic=True)

    ws["A7"] = "Initial deposit (£)"
    ws["B7"] = scenario["deposit"]
    ws["B7"].number_format = "£#,##0.00"

    ws["A8"] = (
        f"Monthly effective rate of interest from {_date_str(scenario['deposit_date'])} "
        f"to {_date_str(scenario['switch_date_last_day'])}"
    )
    ws["B8"] = scenario["initial_monthly_rate"] / 100
    ws["B8"].number_format = "0.00%"

    ws["A9"] = f"Annual effective rate of interest from {_date_str(scenario['switch_date'])}"
    ws["B9"] = scenario["annual_rate"] / 100
    ws["B9"].number_format = "0.00%"

    ws["A10"] = f"Monthly-equivalent rate of interest from {_date_str(scenario['switch_date'])}"
    ws["B10"].fill = _ANSWER_FILL
    ws["B10"].number_format = "0.0000%"

    ws["A11"] = "Regular monthly payment (£)"
    ws["B11"] = scenario["payment"]
    ws["B11"].number_format = "£#,##0.00"

    header_row = 13
    headers = ["Date", "Balance before payment (£)", "Payment (£)", "Balance after payment (£)"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    deposit_row = header_row + 1
    ws.cell(row=deposit_row, column=1, value=scenario["deposit_date"]).number_format = "dd/mm/yyyy"
    ws.cell(row=deposit_row, column=2, value=0).number_format = "£#,##0.00"
    ws.cell(row=deposit_row, column=3, value=scenario["deposit"]).number_format = "£#,##0.00"
    ws.cell(row=deposit_row, column=4, value=scenario["deposit"]).number_format = "£#,##0.00"

    first_answer_row = deposit_row + 1
    for i, pay_date in enumerate(scenario["payment_dates"]):
        row = first_answer_row + i
        ws.cell(row=row, column=1, value=pay_date).number_format = "dd/mm/yyyy"
        for col in (2, 3, 4):
            cell = ws.cell(row=row, column=col)
            cell.fill = _ANSWER_FILL
            cell.number_format = "£#,##0.00"

    target_row = first_answer_row + len(scenario["payment_dates"]) - 1
    target_cell = f"{get_column_letter(2)}{target_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), target_cell


def _generate_scenario():
    """Builds one randomised savings-schedule scenario: the input parameters, the full
    computed month-by-month schedule, and the question text/scaffold/worked-solution derived
    from it. Split out from `generate_savings_schedule_question()` so pipeline scripts (e.g.
    building a multi-question homework spreadsheet) can reuse the same scenario/schedule
    generation without duplicating the recurrence logic."""
    name = random.choice(_NAMES)
    pronoun, possessive = _PRONOUNS[name]
    item = random.choice(_SCENARIOS)

    deposit_year = random.randint(2023, 2025)
    deposit_month = random.randint(1, 12)
    deposit_mi = _ym_to_mi(deposit_year, deposit_month)
    deposit_date = date(deposit_year, deposit_month, 1)

    switch_after_months = random.randint(3, 6)
    switch_date = _add_months(deposit_mi, switch_after_months)
    switch_date_last_day = _last_day_of_month(_add_months(deposit_mi, switch_after_months - 1))

    deposit = random.choice(_DEPOSITS)
    payment = random.choice(_PAYMENTS)
    initial_monthly_rate = random.choice(_INITIAL_MONTHLY_RATES)
    annual_rate = random.choice(_ANNUAL_RATES)
    monthly_equiv_rate = (1 + annual_rate / 100) ** (1 / 12) - 1

    payment_dates = [_add_months(deposit_mi, k) for k in range(1, 13)]
    target_date = payment_dates[-1]

    scenario = {
        "name": name,
        "deposit": deposit,
        "payment": payment,
        "deposit_date": deposit_date,
        "switch_date": switch_date,
        "switch_date_last_day": switch_date_last_day,
        "initial_monthly_rate": initial_monthly_rate,
        "annual_rate": annual_rate,
        "monthly_equiv_rate": monthly_equiv_rate,
        "payment_dates": payment_dates,
    }

    balance = float(deposit)
    worked = [
        f"Monthly-equivalent of {annual_rate}% per year from {_date_str(switch_date)} "
        f"= (1 + {annual_rate / 100:.4f})^(1/12) − 1 = {monthly_equiv_rate * 100:.4f}% per month",
        f"{_date_str(deposit_date)}: opening deposit = £{_fmt_money(balance)}",
    ]
    schedule_rows = []
    for pay_date in payment_dates:
        rate = initial_monthly_rate / 100 if pay_date < switch_date else monthly_equiv_rate
        interest = _excel_round(rate * balance, 2)
        before = _excel_round(balance + interest, 2)
        after = _excel_round(before + payment, 2)
        worked.append(
            f"{_date_str(pay_date)}: £{_fmt_money(balance)} + £{_fmt_money(interest)} interest "
            f"= £{_fmt_money(before)}; + £{payment} payment = £{_fmt_money(after)}"
        )
        schedule_rows.append((pay_date, before, payment, after))
        balance = after
        target_before = before

    scenario["schedule_rows"] = schedule_rows
    answer = round(target_before, 2)

    question_text = (
        f"{name} has been saving for {item}. {pronoun.capitalize()} opened a savings account "
        f"with an initial deposit of £{deposit} on {_date_str(deposit_date)}. "
        f"{pronoun.capitalize()} made regular monthly payments of £{payment} into the account "
        f"on the first day of each month from {_date_str(payment_dates[0])}.\n\n"
        f"The effective rates of interest for the savings account are as follows:\n"
        f"- {_date_str(deposit_date)} to {_date_str(switch_date_last_day)}: {initial_monthly_rate}% per month\n"
        f"- From {_date_str(switch_date)}: {annual_rate}% per year\n\n"
        f"Download the spreadsheet below and use it to calculate the balance in {possessive} "
        f"account immediately before {pronoun} makes {possessive} payment on {_date_str(target_date)}."
    )

    scaffold_steps = [
        {
            "prompt": f"What is the monthly-equivalent of the annual rate that applies "
                      f"from {_date_str(switch_date)}? (as a percentage, 4 d.p.)",
            "answer": round(monthly_equiv_rate * 100, 4),
        },
        {
            "prompt": f"What is the balance immediately before the payment on {_date_str(target_date)}?",
            "answer": answer,
        },
    ]

    return {
        "scenario": scenario,
        "name": name,
        "question_text": question_text,
        "scaffold_steps": scaffold_steps,
        "worked_solution": worked,
        "answer": answer,
    }


def generate_savings_schedule_question():
    gen = _generate_scenario()
    spreadsheet_bytes, target_cell = _build_workbook(gen["scenario"])

    return Question(
        question_text=gen["question_text"],
        correct_answer=gen["answer"],
        topic="Finance",
        question_type="Savings Schedule",
        scaffold_steps=gen["scaffold_steps"],
        worked_solution=gen["worked_solution"],
        notes=NOTES,
        metadata={
            "spreadsheet_bytes": spreadsheet_bytes,
            "spreadsheet_filename": f"savings_schedule_{gen['name'].lower()}.xlsx",
            "spreadsheet_answer_cell": ("Savings", target_cell),
        },
    )
